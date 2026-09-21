"""Rejected-row export.

A rejection the operator cannot act on is only half a message. When rows are
dropped because a branch or product is not in the master, the fix is to add the
master record and load those rows -- so the platform hands back a workbook
containing exactly the dropped rows, in the original column layout, with the
reason attached.

That file is re-uploadable as-is once the master data is corrected: the reason
columns sit to the right of the mapped range and are ignored on the way back in.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import StagingAccountData, UploadBatch, UploadException

#: The mapped input columns, in the order the bank's file carries them. Keeping
#: the layout identical is what makes the download re-uploadable without edits.
SOURCE_COLUMNS = [
    ("Branch ", "branch_code"),
    ("acno", "account_no"),
    ("Type", "side"),
    ("balance", "balance"),
    ("inttpbl", "int_payable"),
    ("inttrcvbl", "int_receivable"),
    ("Prdct", "product_code"),
    ("ROI", "roi"),
]
#: Diagnostics, appended to the right of the mapped range.
REASON_COLUMNS = ["Source row", "Rule", "Reason"]

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_REASON_FILL = PatternFill("solid", fgColor="C00000")


def build_rejects_workbook(session: Session, batch: UploadBatch) -> tuple[bytes, int]:
    """Return (xlsx bytes, row count) for the rows this batch could not accept."""
    rows = list(session.scalars(
        select(StagingAccountData)
        .where(StagingAccountData.batch_id == batch.id)
        .where(StagingAccountData.status == "REJECT")
        .order_by(StagingAccountData.source_row_no)
    ))

    reasons: dict[int, list[UploadException]] = defaultdict(list)
    for exc in session.scalars(
        select(UploadException)
        .where(UploadException.batch_id == batch.id)
        .where(UploadException.severity == "REJECT")
        .order_by(UploadException.source_row_no)
    ):
        if exc.source_row_no:
            reasons[exc.source_row_no].append(exc)

    wb = Workbook()

    # --- sheet 1: the rows themselves, ready to re-upload ------------------
    ws = wb.active
    ws.title = _sheet_title(batch)

    headers = [h for h, _ in SOURCE_COLUMNS] + REASON_COLUMNS
    ws.append(headers)
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _REASON_FILL if i > len(SOURCE_COLUMNS) else _HEADER_FILL
        c.alignment = Alignment(horizontal="center")

    for r in rows:
        found = reasons.get(r.source_row_no or 0, [])
        ws.append([
            *[getattr(r, attr) for _, attr in SOURCE_COLUMNS],
            r.origin or r.source_row_no,
            ", ".join(sorted({e.rule_code for e in found})),
            " | ".join(e.message for e in found) or "rejected",
        ])

    widths = [10, 16, 7, 16, 14, 14, 14, 9, 16, 10, 74]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    # --- sheet 2: what to fix, grouped so the work is obvious --------------
    summary = wb.create_sheet("What to fix")
    summary.append(["Batch", batch.batch_ref])
    summary.append(["File", batch.file_name])
    summary.append(["Business date",
                    batch.business_date.isoformat() if batch.business_date else ""])
    summary.append(["Rows read", batch.total_rows])
    summary.append(["Rows accepted", batch.accepted_rows])
    summary.append(["Rows rejected", len(rows)])
    summary.append(["Generated", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")])
    summary.append([])

    summary.append(["Unregistered branches", "Rows affected"])
    _bold_row(summary)
    missing_branches = _count_missing(rows, reasons, "V003", "branch_code")
    for code, count in sorted(missing_branches.items()):
        summary.append([code, count])
    if not missing_branches:
        summary.append(["(none)", 0])

    summary.append([])
    summary.append(["Unregistered or mismatched products", "Rows affected"])
    _bold_row(summary)
    missing_products = _count_missing(rows, reasons, ("V004", "V005"), "product_code")
    for code, count in sorted(missing_products.items()):
        summary.append([code, count])
    if not missing_products:
        summary.append(["(none)", 0])

    summary.append([])
    summary.append(["Every rule that fired", "Rows affected"])
    _bold_row(summary)
    by_rule: dict[str, int] = defaultdict(int)
    for found in reasons.values():
        for code in {e.rule_code for e in found}:
            by_rule[code] += 1
    for code, count in sorted(by_rule.items()):
        summary.append([code, count])

    summary.append([])
    summary.append(["Next step:"])
    summary.append(["1. Add the branches and products listed above to the master data."])
    summary.append(["2. Re-upload THIS file. The reason columns are ignored on re-import."])
    summary.append(["3. These rows will merge with the day already loaded."])

    summary.column_dimensions["A"].width = 52
    summary.column_dimensions["B"].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def _bold_row(ws) -> None:
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


def _count_missing(rows, reasons, codes, attr) -> dict[str, int]:
    wanted = {codes} if isinstance(codes, str) else set(codes)
    out: dict[str, int] = defaultdict(int)
    for r in rows:
        fired = {e.rule_code for e in reasons.get(r.source_row_no or 0, [])}
        if fired & wanted:
            out[str(getattr(r, attr) or "(blank)")] += 1
    return dict(out)


def _sheet_title(batch: UploadBatch) -> str:
    """Name the sheet the way a daily file is named.

    The point of this workbook is that it can be handed straight back once the
    master data is fixed, so its data sheet has to look like any other daily
    sheet to the importer. A descriptive title such as "Rejected 2026-09-05"
    would not parse as a date and the file would be unimportable -- which would
    defeat the whole feature.
    """
    if batch.business_date:
        # "%-d" is not portable; strip the leading zero explicitly.
        return batch.business_date.strftime("%d %b %y").lstrip("0")
    return "Rejected rows"
