"""Excel adapter.

Reads .xlsx and .xlsm in openpyxl's read-only streaming mode, so memory is
O(row) rather than O(file).

Security note
-------------
`keep_vba=False` means the VBA project is never loaded, and openpyxl has no
mechanism to execute it in any case. An uploaded .xlsm is treated purely as a
zip of XML. Per plan §15.14 this still runs in a sandboxed worker with no
network egress and a wall-clock cap.

Capacity note
-------------
A worksheet holds at most 1,048,576 rows. Past roughly 1M accounts/day the
format cannot carry the daily feed at all, which is why the adapter boundary
exists -- see plan §15.2.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from app.domain.errors import ValidationError
from app.domain.types import RawRow
from app.ingestion.base import ExtractStats, ProbeResult, SheetInfo, SourceAdapter
from app.ingestion.mapping import ColumnMapping

#: Hard ceiling of the .xlsx format itself.
EXCEL_MAX_ROWS = 1_048_576

#: Header text expected above each mapped column, used for drift detection only.
EXPECTED_HEADERS: dict[str, tuple[str, ...]] = {
    "branch_code": ("branch",),
    "account_no": ("acno", "account", "account no", "accountno"),
    "side": ("type", "side"),
    "balance": ("balance", "bal"),
    "int_payable": ("inttpbl", "interest payable", "int payable"),
    "int_receivable": ("inttrcvbl", "interest receivable", "int receivable"),
    "product_code": ("prdct", "product", "product code"),
    "roi": ("roi", "rate of interest"),
}


class ExcelAdapter(SourceAdapter):
    """Streams account rows out of a workbook."""

    code = "EXCEL"

    def __init__(
        self,
        mapping: ColumnMapping,
        *,
        business_date: date | None = None,
        sheet_name: str | None = None,
    ) -> None:
        """
        Args:
            mapping: The column profile to read with.
            business_date: The date this file is for. Supply it for the normal
                daily upload, where the sheet may be named anything.
            sheet_name: Read only this worksheet. Needed only to disambiguate a
                multi-sheet workbook.

        Three supported shapes, in order of how often they occur:

        * **Daily upload** -- one sheet, one day. If the sheet name parses as a
          date it is used; otherwise supply `business_date`.
        * **Targeted upload** -- a multi-sheet workbook plus `sheet_name`.
        * **Historical backfill** -- neither argument. Every date-named sheet is
          read and dated from its own name, which is how the legacy workbook's
          six days load in one pass.
        """
        self.mapping = mapping
        self.business_date = business_date
        self.sheet_name = sheet_name

    # ------------------------------------------------------------------ #
    # Probe
    # ------------------------------------------------------------------ #

    def probe(self, source: Path) -> ProbeResult:
        """Inspect without ingesting: which sheets qualify, and why not.

        Sheets that do not qualify are *listed with a reason*, never silently
        passed over. This is what replaces the macro's "any sheet Excel can read
        as a date is input".
        """
        result = ProbeResult(adapter=self.code)

        try:
            wb = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
        except Exception as exc:  # noqa: BLE001 - surfaced to the user verbatim
            result.errors.append(f"cannot open workbook: {exc}")
            return result

        try:
            included, excluded, errors = self._resolve_sheets(list(wb.sheetnames))
            result.errors.extend(errors)
            result.sheets.extend(excluded)

            for name, bdate in included:
                ws = wb[name]
                info = SheetInfo(name, bdate, True, "included")
                info.data_rows = self._count_data_rows(ws)
                if info.data_rows == 0:
                    info.included = False
                    info.reason = "no data rows below the header"
                result.sheets.append(info)

                if info.included:
                    result.header_issues.extend(self._check_headers(ws, name))

            # Keep workbook order so the preview reads like the file looks.
            order = {n: i for i, n in enumerate(wb.sheetnames)}
            result.sheets.sort(key=lambda s: order.get(s.name, 0))
        finally:
            wb.close()

        return result

    # ------------------------------------------------------------------ #
    # Extract
    # ------------------------------------------------------------------ #

    def extract(
        self, source: Path, stats: ExtractStats | None = None
    ) -> Iterator[RawRow]:
        """Yield one `RawRow` per data row, across every qualifying sheet.

        `source_row_no` is a running counter over the whole file so it is unique
        within a batch; `extras["_origin"]` carries the human-readable
        "<sheet>!<row>" so an exception points at a cell the user can open.
        """
        stats = stats if stats is not None else ExtractStats()
        wb = load_workbook(source, read_only=True, data_only=True, keep_vba=False)
        counter = 0

        try:
            included, _excluded, errors = self._resolve_sheets(list(wb.sheetnames))
            if errors:
                raise ValidationError("V001", "; ".join(errors), field="sheet")

            for name, bdate in included:
                ws = wb[name]
                for excel_row, values in enumerate(
                    ws.iter_rows(min_row=self.mapping.first_data_row, values_only=True),
                    start=self.mapping.first_data_row,
                ):
                    stats.physical_rows += 1

                    if values is None or all(v is None for v in values):
                        stats.blank_rows += 1
                        continue

                    if self._is_structural(values):
                        stats.structural_rows += 1
                        continue

                    counter += 1
                    stats.data_rows += 1
                    yield self._to_row(values, counter, bdate, name, excel_row)
        finally:
            wb.close()

    # ------------------------------------------------------------------ #
    # Internals
    # ------------------------------------------------------------------ #

    def _resolve_sheets(
        self, sheetnames: list[str]
    ) -> tuple[list[tuple[str, date]], list[SheetInfo], list[str]]:
        """Decide which worksheets are read and what date each one carries.

        Returns `(included, excluded, errors)`. Every sheet lands in `included`
        or `excluded`; an excluded sheet always carries a reason, because a
        silently ignored sheet is how a day goes missing unnoticed.
        """
        included: list[tuple[str, date]] = []
        excluded: list[SheetInfo] = []
        errors: list[str] = []

        # --- targeted: one named sheet ---------------------------------- #
        if self.sheet_name is not None:
            if self.sheet_name not in sheetnames:
                errors.append(
                    f"sheet {self.sheet_name!r} not found; workbook has {sheetnames}"
                )
                return included, excluded, errors
            bdate = self.business_date or self.mapping.sheet_selector.parse(self.sheet_name)
            if bdate is None:
                errors.append(
                    f"cannot determine a business date for sheet {self.sheet_name!r}; "
                    "supply business_date with the upload"
                )
                return included, excluded, errors
            included.append((self.sheet_name, bdate))
            excluded.extend(
                SheetInfo(n, None, False, "not the selected sheet")
                for n in sheetnames
                if n != self.sheet_name
            )
            return included, excluded, errors

        # --- daily: an explicit date, so the file must be unambiguous ---- #
        if self.business_date is not None:
            dated = [n for n in sheetnames if self.mapping.sheet_selector.parse(n)]
            if len(sheetnames) == 1:
                included.append((sheetnames[0], self.business_date))
            elif len(dated) == 1:
                included.append((dated[0], self.business_date))
                excluded.extend(
                    SheetInfo(n, None, False, "sheet name is not a business date")
                    for n in sheetnames
                    if n != dated[0]
                )
            else:
                errors.append(
                    f"business_date was supplied but the workbook has "
                    f"{len(sheetnames)} worksheets ({len(dated)} date-named); "
                    "specify sheet_name, or omit business_date to load every "
                    "date-named sheet"
                )
            return included, excluded, errors

        # --- backfill: every date-named sheet, dated from its own name --- #
        for name in sheetnames:
            bdate = self.mapping.sheet_selector.parse(name)
            if bdate is None:
                excluded.append(
                    SheetInfo(name, None, False, "sheet name is not a business date")
                )
            else:
                included.append((name, bdate))

        if not included:
            hint = (
                f"no worksheet name parsed as a date (tried "
                f"{list(self.mapping.sheet_selector.patterns)})."
            )
            if len(sheetnames) == 1:
                hint += (
                    f" The workbook has a single sheet {sheetnames[0]!r} -- supply "
                    "business_date with the upload to load it as a daily file."
                )
            errors.append(hint)

        return included, excluded, errors

    def _is_structural(self, values: tuple[Any, ...]) -> bool:
        """A declared non-data row, e.g. a branch subtotal.

        The legacy macro excluded these only because column A happened to be
        blank. Here the rule is explicit, so a subtotal row that acquires a
        label cannot become a phantom account (defect L4).
        """
        for field_name in self.mapping.ignore_row_when_blank:
            spec = self.mapping.columns.get(field_name)
            if spec is None:
                continue
            value = self._cell(values, spec.index)
            if value is None or (isinstance(value, str) and not value.strip()):
                return True
        return False

    @staticmethod
    def _cell(values: tuple[Any, ...], index: int) -> Any:
        return values[index] if 0 <= index < len(values) else None

    def _to_row(
        self,
        values: tuple[Any, ...],
        counter: int,
        bdate: date,
        sheet: str,
        excel_row: int,
    ) -> RawRow:
        fields: dict[str, Any] = {}
        errors: list[str] = []

        for field_name in self.mapping.columns:
            raw = self._cell(values, self.mapping.columns[field_name].index)
            try:
                fields[field_name] = self.mapping.coerce(field_name, raw)
            except ValidationError as exc:
                # Coercion failures are data, not crashes: the row is still
                # emitted so validation can reject it with a precise message
                # and the user can see every problem in one pass.
                fields[field_name] = None
                errors.append(f"{exc.code}:{exc.message}")

        extras: dict[str, Any] = {"_origin": f"{sheet}!{excel_row}", "_sheet": sheet}
        for letter, key in self.mapping.capture_as_extras.items():
            idx = ord(letter.upper()) - 65
            extras[key] = self._cell(values, idx)
        if errors:
            extras["_coercion_errors"] = errors

        return RawRow(
            source_row_no=counter,
            business_date=bdate,
            branch_code=fields.get("branch_code"),
            account_no=fields.get("account_no"),
            side=fields.get("side"),
            product_code=fields.get("product_code"),
            balance=fields.get("balance"),
            int_payable=fields.get("int_payable"),
            int_receivable=fields.get("int_receivable"),
            roi=fields.get("roi"),
            extras=extras,
        )

    def _count_data_rows(self, ws: Any) -> int:
        count = 0
        for values in ws.iter_rows(min_row=self.mapping.first_data_row, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            if self._is_structural(values):
                continue
            count += 1
        return count

    def _check_headers(self, ws: Any, sheet: str) -> list[str]:
        """Warn when a header does not look like what the mapping expects.

        Advisory only -- position is authoritative. A bank that renames a header
        should not have its file rejected, but somebody should be told.
        """
        issues: list[str] = []
        header = next(
            ws.iter_rows(
                min_row=self.mapping.header_row,
                max_row=self.mapping.header_row,
                values_only=True,
            ),
            None,
        )
        if header is None:
            return [f"{sheet}: no header row at row {self.mapping.header_row}"]

        for field_name, spec in self.mapping.columns.items():
            actual = self._cell(header, spec.index)
            expected = EXPECTED_HEADERS.get(field_name)
            if not expected or actual is None:
                continue
            if str(actual).strip().lower() not in expected:
                issues.append(
                    f"{sheet}: column {spec.letter} header is {str(actual).strip()!r}, "
                    f"expected one of {expected} for {field_name}"
                )
        return issues
